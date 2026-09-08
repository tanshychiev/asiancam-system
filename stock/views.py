from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.urls import reverse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from openpyxl import Workbook, load_workbook

from accounting.models import JournalEntry, JournalEntryLine
from core.models import Company

from .forms import (
    ItemBrandForm,
    ItemForm,
    ItemGroupForm,
    StockDocumentForm,
    StockDocumentLineFormSet,
    UnitSetForm,
    WarehouseForm,
)
from .models import (
    Item,
    ItemBrand,
    ItemGroup,
    StockDocument,
    UnitSet,
    Warehouse,
)


# =========================================================
# HELPERS
# =========================================================

def get_selected_company(request):
    company_id = request.session.get("selected_company_id")

    if not company_id:
        return None

    return Company.objects.filter(
        id=company_id,
        is_active=True,
    ).first()


def can_access_company(user, company):
    if not company:
        return False

    if user.is_superuser:
        return True

    profile = getattr(user, "profile", None)

    if profile and getattr(profile, "user_type", None) == "client":
        return profile.company_id == company.id

    if hasattr(company, "assigned_staff"):
        return company.assigned_staff.filter(id=user.id).exists()

    return False


def require_company_access(request):
    company = get_selected_company(request)

    if not company:
        messages.warning(request, "Please select a company first.")
        return None, redirect("company_list")

    if not can_access_company(request.user, company):
        messages.error(request, "You do not have permission to access this company.")
        return None, redirect("company_list")

    return company, None


def get_posted_status():
    return getattr(JournalEntry, "STATUS_POSTED", "posted")


def create_stock_journal(stock_doc, user):
    """
    Auto journal logic.

    Stock Issue example:
        Dr COGS / Expense
        Cr Inventory Asset

    Stock Adjustment example:
        Dr Inventory / Adjustment Loss
        Cr Inventory / Adjustment Gain

    Stock Transfer:
        No journal because it only moves stock between warehouses.

    Stock Assembly:
        Can create journal if debit/credit accounts are selected.
    """

    if stock_doc.status != StockDocument.STATUS_POSTED:
        return None

    if stock_doc.document_type == StockDocument.TYPE_TRANSFER:
        return None

    if not stock_doc.debit_account or not stock_doc.credit_account:
        return None

    total_amount = stock_doc.total_amount

    if total_amount <= 0:
        return None

    with transaction.atomic():
        old_entry = stock_doc.journal_entry

        if old_entry:
            old_entry.delete()

        entry = JournalEntry.objects.create(
            company=stock_doc.company,
            entry_date=stock_doc.document_date,
            reference_no=stock_doc.number,
            description=f"{stock_doc.get_document_type_display()} - {stock_doc.memo}",
            status=get_posted_status(),
            created_by=user,
        )

        JournalEntryLine.objects.create(
            journal_entry=entry,
            account=stock_doc.debit_account,
            description=stock_doc.memo or stock_doc.get_document_type_display(),
            debit=total_amount,
            credit=Decimal("0.00"),
        )

        JournalEntryLine.objects.create(
            journal_entry=entry,
            account=stock_doc.credit_account,
            description=stock_doc.memo or stock_doc.get_document_type_display(),
            debit=Decimal("0.00"),
            credit=total_amount,
        )

        stock_doc.journal_entry = entry
        stock_doc.save(update_fields=["journal_entry"])

        return entry


# =========================================================
# ITEM MASTER
# =========================================================

def _item_has_field(name):
    try:
        Item._meta.get_field(name)
        return True
    except Exception:
        return False


def _safe_decimal(value, default=None):
    from decimal import Decimal, InvalidOperation
    if value is None or str(value).strip() == "":
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError, TypeError):
        return default


def _item_filter_queryset(request, company):
    query = (request.GET.get("q") or "").strip()
    item_type = (request.GET.get("item_type") or "").strip()
    group_id = (request.GET.get("group") or "").strip()
    brand_id = (request.GET.get("brand") or "").strip()
    unit_id = (request.GET.get("unit") or "").strip()
    active = (request.GET.get("active") or "").strip()

    items = (
        Item.objects
        .filter(company=company)
        .select_related("item_group", "item_brand", "unit_set")
    )

    if query:
        search = (
            Q(code__icontains=query)
            | Q(name__icontains=query)
            | Q(memo__icontains=query)
        )
        if _item_has_field("description"):
            search |= Q(description__icontains=query)
        if _item_has_field("barcode"):
            search |= Q(barcode__icontains=query)
        items = items.filter(search)

    if item_type:
        items = items.filter(item_type=item_type)
    if group_id:
        items = items.filter(item_group_id=group_id)
    if brand_id:
        items = items.filter(item_brand_id=brand_id)
    if unit_id:
        items = items.filter(unit_set_id=unit_id)
    if active == "1":
        items = items.filter(is_active=True)
    elif active == "0":
        items = items.filter(is_active=False)

    sort = (request.GET.get("sort") or "name").strip()
    sort_map = {
        "name": "name",
        "-name": "-name",
        "code": "code",
        "-code": "-code",
        "cost": "cost_price",
        "-cost": "-cost_price",
        "sale": "sale_price",
        "-sale": "-sale_price",
    }
    items = items.order_by(sort_map.get(sort, "name"), "code")

    return items, {
        "query": query,
        "item_type": item_type,
        "group_id": group_id,
        "brand_id": brand_id,
        "unit_id": unit_id,
        "active": active,
        "sort": sort,
    }


@login_required
def item_list(request):
    company, response = require_company_access(request)
    if response:
        return response

    # Toolbar / row actions all post back to this one page.
    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()
        selected_ids = [
            int(x) for x in request.POST.getlist("selected")
            if str(x).isdigit()
        ]
        selected = Item.objects.filter(company=company, id__in=selected_ids)

        if action == "delete":
            if not selected_ids:
                messages.warning(request, "Select at least one item.")
            else:
                deleted = 0
                protected = 0
                for item in selected:
                    try:
                        item.delete()
                        deleted += 1
                    except Exception:
                        # Safer than breaking accounting/stock history.
                        item.is_active = False
                        item.save(update_fields=["is_active"])
                        protected += 1
                if deleted:
                    messages.success(request, f"{deleted} item(s) deleted.")
                if protected:
                    messages.warning(
                        request,
                        f"{protected} item(s) had history and were deactivated instead.",
                    )
            return redirect(request.get_full_path())

        if action == "bulk_update":
            if not selected_ids:
                messages.warning(request, "Select at least one item.")
                return redirect(request.get_full_path())

            changed = 0
            field = (request.POST.get("bulk_field") or "").strip()
            value = (request.POST.get("bulk_value") or "").strip()

            allowed = {"is_active", "item_type", "item_group", "item_brand", "unit_set", "cost_price", "sale_price"}
            if field not in allowed:
                messages.error(request, "Invalid bulk update field.")
                return redirect(request.get_full_path())

            for item in selected:
                if field == "is_active":
                    setattr(item, field, value == "1")
                elif field in {"item_group", "item_brand", "unit_set"}:
                    setattr(item, f"{field}_id", int(value) if value.isdigit() else None)
                elif field in {"cost_price", "sale_price"}:
                    amount = _safe_decimal(value)
                    if amount is None or amount < 0:
                        continue
                    setattr(item, field, amount)
                else:
                    valid_types = {x[0] for x in Item.ITEM_TYPE_CHOICES}
                    if value not in valid_types:
                        continue
                    item.item_type = value
                item.save()
                changed += 1

            messages.success(request, f"{changed} item(s) updated.")
            return redirect(request.get_full_path())

        if action == "price_update":
            item_id = request.POST.get("item_id")
            item = get_object_or_404(Item, company=company, id=item_id)
            cost = _safe_decimal(request.POST.get("cost_price"))
            sale = _safe_decimal(request.POST.get("sale_price"))
            if cost is not None and cost >= 0:
                item.cost_price = cost
            if sale is not None and sale >= 0:
                item.sale_price = sale
            update_fields = ["cost_price", "sale_price"]

            if _item_has_field("barcode"):
                setattr(item, "barcode", (request.POST.get("barcode") or "").strip())
                update_fields.append("barcode")
            if _item_has_field("base_price"):
                base_price = _safe_decimal(request.POST.get("base_price"))
                if base_price is not None and base_price >= 0:
                    setattr(item, "base_price", base_price)
                    update_fields.append("base_price")

            item.save(update_fields=list(dict.fromkeys(update_fields)))
            messages.success(request, f"{item.name} price information updated.")
            return redirect(request.get_full_path())

        if action == "duplicate":
            item_id = request.POST.get("item_id")
            source = get_object_or_404(Item, company=company, id=item_id)

            original_code = source.code or "ITEM"
            suffix = 1
            new_code = f"{original_code}-COPY"
            while Item.objects.filter(company=company, code=new_code).exists():
                suffix += 1
                new_code = f"{original_code}-COPY{suffix}"

            clone = Item()
            for field in Item._meta.concrete_fields:
                if field.primary_key or field.name in {"id", "company"}:
                    continue
                if field.name == "code":
                    setattr(clone, field.name, new_code)
                elif field.name == "name":
                    setattr(clone, field.name, f"{source.name} Copy")
                else:
                    setattr(clone, field.attname, getattr(source, field.attname))
            clone.company = company
            clone.save()
            messages.success(request, f"Duplicated as {clone.code} - {clone.name}.")
            return redirect("stock_item_edit", item_id=clone.id)

    items, filters = _item_filter_queryset(request, company)

    try:
        per_page = int(request.GET.get("per_page") or 10)
    except (TypeError, ValueError):
        per_page = 10
    if per_page not in {10, 25, 50, 100}:
        per_page = 10

    paginator = Paginator(items, per_page)
    page_obj = paginator.get_page(request.GET.get("page") or 1)

    groups = ItemGroup.objects.filter(company=company, is_active=True).order_by("name")
    brands = ItemBrand.objects.filter(company=company, is_active=True).order_by("name")
    units = UnitSet.objects.filter(company=company, is_active=True).order_by("name")

    context = {
        "company": company,
        "items": page_obj.object_list,
        "page_obj": page_obj,
        "paginator": paginator,
        "per_page": per_page,
        "item_type_choices": Item.ITEM_TYPE_CHOICES,
        "groups": groups,
        "brands": brands,
        "units": units,
        "has_barcode": _item_has_field("barcode"),
        "has_description": _item_has_field("description"),
        "has_cost_method": _item_has_field("cost_method"),
        "has_base_price": _item_has_field("base_price"),
        **filters,
    }
    return render(request, "stock/item_list.html", context)


@login_required
def item_delete(request, item_id):
    company, response = require_company_access(request)
    if response:
        return response
    item = get_object_or_404(Item, id=item_id, company=company)
    if request.method != "POST":
        return redirect("stock_item_list")
    try:
        name = item.name
        item.delete()
        messages.success(request, f"{name} deleted.")
    except Exception:
        item.is_active = False
        item.save(update_fields=["is_active"])
        messages.warning(
            request,
            f"{item.name} has transaction history, so it was deactivated instead of deleted.",
        )
    return redirect("stock_item_list")


@login_required
def item_duplicate(request, item_id):
    company, response = require_company_access(request)
    if response:
        return response
    if request.method != "POST":
        return redirect("stock_item_list")

    source = get_object_or_404(Item, id=item_id, company=company)
    base = source.code or "ITEM"
    code = f"{base}-COPY"
    n = 2
    while Item.objects.filter(company=company, code=code).exists():
        code = f"{base}-COPY{n}"
        n += 1

    clone = Item()
    for field in Item._meta.concrete_fields:
        if field.primary_key or field.name in {"id", "company"}:
            continue
        if field.name == "code":
            setattr(clone, field.name, code)
        elif field.name == "name":
            setattr(clone, field.name, f"{source.name} Copy")
        else:
            setattr(clone, field.attname, getattr(source, field.attname))
    clone.company = company
    clone.save()
    messages.success(request, f"Item duplicated as {clone.code}.")
    return redirect("stock_item_edit", item_id=clone.id)


@login_required
def item_export_excel(request):
    company, response = require_company_access(request)
    if response:
        return response

    items, _ = _item_filter_queryset(request, company)

    wb = Workbook()
    ws = wb.active
    ws.title = "Items"

    headers = [
        "ITEM_CODE", "ITEM_NAME", "ITEM_TYPE", "ITEM_GROUP", "ITEM_BRAND",
        "UNIT_SET", "COST_PRICE", "SALE_PRICE", "ACTIVE", "MEMO"
    ]
    optional = []
    if _item_has_field("barcode"):
        optional.append(("BARCODE", "barcode"))
    if _item_has_field("description"):
        optional.append(("DESCRIPTION", "description"))
    if _item_has_field("cost_method"):
        optional.append(("COST_METHOD", "cost_method"))
    if _item_has_field("base_price"):
        optional.append(("BASE_PRICE", "base_price"))

    # Keep commonly requested optional columns close to Code/Name.
    headers = ["ITEM_CODE"] + [h for h, _ in optional if h == "BARCODE"] + ["ITEM_NAME"] + \
              [h for h, _ in optional if h in {"DESCRIPTION", "COST_METHOD"}] + \
              ["ITEM_TYPE", "ITEM_GROUP", "ITEM_BRAND", "UNIT_SET", "COST_PRICE"] + \
              [h for h, _ in optional if h == "BASE_PRICE"] + ["SALE_PRICE", "ACTIVE", "MEMO"]
    ws.append(headers)

    for item in items:
        data = {
            "ITEM_CODE": item.code,
            "ITEM_NAME": item.name,
            "ITEM_TYPE": item.item_type,
            "ITEM_GROUP": item.item_group.name if item.item_group else "",
            "ITEM_BRAND": item.item_brand.name if item.item_brand else "",
            "UNIT_SET": item.unit_set.name if item.unit_set else "",
            "COST_PRICE": item.cost_price,
            "SALE_PRICE": item.sale_price,
            "ACTIVE": "Yes" if item.is_active else "No",
            "MEMO": item.memo,
            "BARCODE": getattr(item, "barcode", ""),
            "DESCRIPTION": getattr(item, "description", ""),
            "COST_METHOD": getattr(item, "cost_method", ""),
            "BASE_PRICE": getattr(item, "base_price", ""),
        }
        ws.append([data.get(h, "") for h in headers])

    ws.freeze_panes = "A2"
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(
            34,
            max(12, max(len(str(c.value or "")) for c in col) + 2)
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="item_list.xlsx"'
    wb.save(response)
    return response


@login_required
def item_import_excel(request):
    company, response = require_company_access(request)
    if response:
        return response

    errors = []
    if request.method == "POST":
        upload = request.FILES.get("file")
        if not upload:
            errors.append("Choose an Excel file.")
        else:
            try:
                wb = load_workbook(upload, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    errors.append("The Excel file is empty.")
                else:
                    headers = [str(x or "").strip().upper() for x in rows[0]]
                    index = {name: i for i, name in enumerate(headers)}

                    def value(row, name, default=""):
                        i = index.get(name)
                        return default if i is None or i >= len(row) else row[i]

                    success = 0
                    with transaction.atomic():
                        for row_no, row in enumerate(rows[1:], 2):
                            code = str(value(row, "ITEM_CODE") or "").strip()
                            name = str(value(row, "ITEM_NAME") or "").strip()
                            if not code and not name:
                                continue
                            if not code or not name:
                                errors.append(f"Row {row_no}: ITEM_CODE and ITEM_NAME are required.")
                                continue

                            group_name = str(value(row, "ITEM_GROUP") or "").strip()
                            brand_name = str(value(row, "ITEM_BRAND") or "").strip()
                            unit_name = str(value(row, "UNIT_SET") or "").strip()

                            group = None
                            brand = None
                            unit = None
                            if group_name:
                                group, _ = ItemGroup.objects.get_or_create(
                                    company=company, name=group_name,
                                    defaults={"is_active": True},
                                )
                            if brand_name:
                                brand, _ = ItemBrand.objects.get_or_create(
                                    company=company, name=brand_name,
                                    defaults={"is_active": True},
                                )
                            if unit_name:
                                unit, _ = UnitSet.objects.get_or_create(
                                    company=company, name=unit_name,
                                    defaults={
                                        "base_unit": "pcs",
                                        "default_purchase": "pcs",
                                        "default_sale": "pcs",
                                        "is_active": True,
                                    },
                                )

                            item_type = str(value(row, "ITEM_TYPE") or Item.TYPE_STOCK_PART).strip()
                            valid_types = {x[0] for x in Item.ITEM_TYPE_CHOICES}
                            if item_type not in valid_types:
                                # also accept display labels such as "Service"
                                label_map = {label.lower(): key for key, label in Item.ITEM_TYPE_CHOICES}
                                item_type = label_map.get(item_type.lower(), Item.TYPE_STOCK_PART)

                            defaults = {
                                "name": name,
                                "item_type": item_type,
                                "item_group": group,
                                "item_brand": brand,
                                "unit_set": unit,
                                "cost_price": _safe_decimal(value(row, "COST_PRICE"), Decimal("0.00")) or Decimal("0.00"),
                                "sale_price": _safe_decimal(value(row, "SALE_PRICE"), Decimal("0.00")) or Decimal("0.00"),
                                "memo": str(value(row, "MEMO") or "").strip(),
                                "is_active": str(value(row, "ACTIVE") or "Yes").strip().lower() not in {"0","false","no","inactive"},
                            }

                            if _item_has_field("barcode"):
                                defaults["barcode"] = str(value(row, "BARCODE") or "").strip()
                            if _item_has_field("description"):
                                defaults["description"] = str(value(row, "DESCRIPTION") or "").strip()
                            if _item_has_field("cost_method"):
                                cm = str(value(row, "COST_METHOD") or "").strip()
                                if cm:
                                    defaults["cost_method"] = cm
                            if _item_has_field("base_price"):
                                defaults["base_price"] = _safe_decimal(value(row, "BASE_PRICE"), Decimal("0.00")) or Decimal("0.00")

                            Item.objects.update_or_create(
                                company=company,
                                code=code,
                                defaults=defaults,
                            )
                            success += 1

                    if errors:
                        messages.warning(
                            request,
                            f"Imported/updated {success} item(s), with {len(errors)} row issue(s)."
                        )
                    else:
                        messages.success(request, f"Imported/updated {success} item(s).")
                        return redirect("stock_item_list")
            except Exception as exc:
                errors.append(f"Could not import Excel: {exc}")

    return render(request, "stock/item_import.html", {
        "company": company,
        "errors": errors,
        "has_barcode": _item_has_field("barcode"),
        "has_description": _item_has_field("description"),
        "has_cost_method": _item_has_field("cost_method"),
        "has_base_price": _item_has_field("base_price"),
    })


@login_required
def item_import_sample(request):
    company, response = require_company_access(request)
    if response:
        return response

    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    headers = [
        "ITEM_CODE",
        "ITEM_NAME",
        "ITEM_TYPE",
        "ITEM_GROUP",
        "ITEM_BRAND",
        "UNIT_SET",
        "COST_PRICE",
        "SALE_PRICE",
        "ACTIVE",
        "MEMO",
    ]
    if _item_has_field("barcode"):
        headers.insert(1, "BARCODE")
    if _item_has_field("description"):
        headers.insert(headers.index("ITEM_TYPE"), "DESCRIPTION")
    if _item_has_field("cost_method"):
        headers.insert(headers.index("ITEM_TYPE"), "COST_METHOD")
    if _item_has_field("base_price"):
        headers.insert(headers.index("SALE_PRICE"), "BASE_PRICE")

    ws.append(headers)
    sample = {
        "ITEM_CODE": "IT-001",
        "BARCODE": "885001",
        "ITEM_NAME": "Sample Item",
        "DESCRIPTION": "Sample description",
        "COST_METHOD": "average",
        "ITEM_TYPE": Item.TYPE_STOCK_PART,
        "ITEM_GROUP": "General",
        "ITEM_BRAND": "General",
        "UNIT_SET": "PCS",
        "COST_PRICE": 1.00,
        "BASE_PRICE": 1.25,
        "SALE_PRICE": 1.50,
        "ACTIVE": "Yes",
        "MEMO": "Replace or delete this sample row.",
    }
    ws.append([sample.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font = c.font.copy(bold=True)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="item_import_sample.xlsx"'
    wb.save(response)
    return response




@login_required
def item_using_info(request, item_id):
    company, response = require_company_access(request)
    if response:
        return response

    item = get_object_or_404(Item, id=item_id, company=company)
    rows = []

    try:
        for line in item.purchase_bill_lines.select_related("bill__vendor").order_by("-bill__bill_date", "-bill_id"):
            bill = line.bill
            edit_url = ""
            try:
                edit_url = reverse("purchase_bill_edit", args=[bill.id])
            except Exception:
                pass
            rows.append({
                "type": "Purchase",
                "number": bill.number or str(bill.id),
                "date": bill.bill_date.strftime("%d-%m-%Y") if bill.bill_date else "",
                "name": getattr(bill.vendor, "name", "") or str(bill.vendor),
                "qty": str(line.qty or ""),
                "amount": str(line.line_amount or ""),
                "edit_url": edit_url,
            })
    except Exception:
        pass

    try:
        for line in item.sales_invoice_lines.select_related("invoice__customer").order_by("-invoice__invoice_date", "-invoice_id"):
            inv = line.invoice
            rows.append({
                "type": inv.get_document_type_display() if hasattr(inv, "get_document_type_display") else "Sale",
                "number": inv.number or str(inv.id),
                "date": inv.invoice_date.strftime("%d-%m-%Y") if inv.invoice_date else "",
                "name": getattr(inv.customer, "name", "") or str(inv.customer),
                "qty": str(line.qty or ""),
                "amount": str(line.line_amount or ""),
                "edit_url": "",
            })
    except Exception:
        pass

    try:
        for line in item.stock_lines.select_related("document").order_by("-document__document_date", "-document_id"):
            doc = line.document
            rows.append({
                "type": doc.get_document_type_display() if hasattr(doc, "get_document_type_display") else "Stock",
                "number": doc.number or str(doc.id),
                "date": doc.document_date.strftime("%d-%m-%Y") if doc.document_date else "",
                "name": str(doc.warehouse or doc.from_warehouse or doc.to_warehouse or ""),
                "qty": str(line.qty or ""),
                "amount": str(line.amount or ""),
                "edit_url": "",
            })
    except Exception:
        pass

    from datetime import datetime
    def sort_key(row):
        try:
            return datetime.strptime(row["date"], "%d-%m-%Y")
        except Exception:
            return datetime.min
    rows.sort(key=sort_key, reverse=True)

    return JsonResponse({
        "item": {"id": item.id, "code": item.code, "name": item.name},
        "count": len(rows),
        "rows": rows,
    })

@login_required
def item_create(request):
    company, response = require_company_access(request)
    if response:
        return response

    if request.method == "POST":
        form = ItemForm(request.POST, company=company)

        if form.is_valid():
            item = form.save(commit=False)
            item.company = company
            item.save()

            messages.success(request, "Item created successfully.")
            return redirect("stock_item_list")
    else:
        form = ItemForm(
            company=company,
            initial={
                "is_active": True,
            },
        )

    return render(request, "stock/item_form.html", {
        "company": company,
        "form": form,
        "page_title": "Create Item",
        "button_text": "Create Item",
    })


@login_required
def item_edit(request, item_id):
    company, response = require_company_access(request)
    if response:
        return response

    item = get_object_or_404(
        Item,
        id=item_id,
        company=company,
    )

    if request.method == "POST":
        form = ItemForm(
            request.POST,
            instance=item,
            company=company,
        )

        if form.is_valid():
            form.save()
            messages.success(request, "Item updated successfully.")
            return redirect("stock_item_list")
    else:
        form = ItemForm(
            instance=item,
            company=company,
        )

    return render(request, "stock/item_form.html", {
        "company": company,
        "form": form,
        "item": item,
        "page_title": "Edit Item",
        "button_text": "Save Changes",
    })


# =========================================================
# MASTER DATA LISTS
# =========================================================

def master_list_view(request, model, template, title):
    company, response = require_company_access(request)
    if response:
        return response

    query = (request.GET.get("q") or "").strip()
    rows = model.objects.filter(company=company)

    if query:
        if model == ItemBrand:
            rows = rows.filter(
                Q(name__icontains=query)
                | Q(description__icontains=query)
            )

        elif model == UnitSet:
            rows = rows.filter(
                Q(name__icontains=query)
                | Q(base_unit__icontains=query)
                | Q(default_purchase__icontains=query)
                | Q(default_sale__icontains=query)
                | Q(memo__icontains=query)
            )

        elif model == Warehouse:
            rows = rows.filter(
                Q(name__icontains=query)
                | Q(code__icontains=query)
                | Q(memo__icontains=query)
            )

        else:
            rows = rows.filter(
                Q(name__icontains=query)
                | Q(code__icontains=query)
                | Q(memo__icontains=query)
            )

    rows = rows.order_by("name")

    return render(request, template, {
        "company": company,
        "rows": rows,
        "query": query,
        "title": title,
    })


@login_required
def item_group_list(request):
    return master_list_view(
        request,
        ItemGroup,
        "stock/item_group_list.html",
        "Item Group",
    )


@login_required
def item_brand_list(request):
    return master_list_view(
        request,
        ItemBrand,
        "stock/item_brand_list.html",
        "Item Brand",
    )


@login_required
def unit_set_list(request):
    return master_list_view(
        request,
        UnitSet,
        "stock/unit_set_list.html",
        "Unit Set",
    )


@login_required
def warehouse_list(request):
    return master_list_view(
        request,
        Warehouse,
        "stock/warehouse_list.html",
        "Warehouse",
    )


# =========================================================
# MASTER DATA CREATE
# =========================================================

def master_create_view(request, form_class, template, title, redirect_name):
    company, response = require_company_access(request)
    if response:
        return response

    if request.method == "POST":
        form = form_class(request.POST)

        if form.is_valid():
            obj = form.save(commit=False)
            obj.company = company
            obj.save()

            messages.success(request, f"{title} saved successfully.")
            return redirect(redirect_name)
    else:
        form = form_class(
            initial={
                "is_active": True,
            },
        )

    return render(request, template, {
        "company": company,
        "form": form,
        "page_title": title,
        "button_text": "Save",
    })


@login_required
def item_group_create(request):
    return master_create_view(
        request,
        ItemGroupForm,
        "stock/master_form.html",
        "Create Item Group",
        "stock_item_group_list",
    )


@login_required
def item_brand_create(request):
    return master_create_view(
        request,
        ItemBrandForm,
        "stock/master_form.html",
        "Create Item Brand",
        "stock_item_brand_list",
    )


@login_required
def unit_set_create(request):
    return master_create_view(
        request,
        UnitSetForm,
        "stock/master_form.html",
        "Create Unit Set",
        "stock_unit_set_list",
    )


@login_required
def warehouse_create(request):
    return master_create_view(
        request,
        WarehouseForm,
        "stock/master_form.html",
        "Create Warehouse",
        "stock_warehouse_list",
    )


# =========================================================
# STOCK DOCUMENT LIST
# =========================================================

@login_required
def stock_document_list(request, doc_type):
    company, response = require_company_access(request)
    if response:
        return response

    query = (request.GET.get("q") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()

    if not date_from and not date_to:
        today = timezone.localdate()
        date_from = today.replace(day=1).strftime("%Y-%m-%d")
        date_to = today.strftime("%Y-%m-%d")

    docs = StockDocument.objects.filter(
        company=company,
        document_type=doc_type,
    )

    if query:
        docs = docs.filter(
            Q(number__icontains=query)
            | Q(memo__icontains=query)
            | Q(lines__item__code__icontains=query)
            | Q(lines__item__name__icontains=query)
        ).distinct()

    if date_from:
        docs = docs.filter(document_date__gte=date_from)

    if date_to:
        docs = docs.filter(document_date__lte=date_to)

    docs = (
        docs
        .select_related(
            "warehouse",
            "from_warehouse",
            "to_warehouse",
            "journal_entry",
        )
        .prefetch_related("lines", "lines__item")
        .order_by("-document_date", "-id")
    )

    title_map = {
        StockDocument.TYPE_ISSUE: "Stock Issue",
        StockDocument.TYPE_ADJUSTMENT: "Stock Adjustment",
        StockDocument.TYPE_ASSEMBLY: "Stock Assembly",
        StockDocument.TYPE_TRANSFER: "Stock Transfer",
    }

    return render(request, "stock/stock_document_list.html", {
        "company": company,
        "docs": docs,
        "doc_type": doc_type,
        "title": title_map.get(doc_type, "Stock Document"),
        "query": query,
        "date_from": date_from,
        "date_to": date_to,
    })


# =========================================================
# STOCK DOCUMENT CREATE
# =========================================================

@login_required
def stock_document_create(request, doc_type):
    company, response = require_company_access(request)
    if response:
        return response

    stock_doc = StockDocument(
        company=company,
        document_type=doc_type,
        created_by=request.user,
    )

    if request.method == "POST":
        form = StockDocumentForm(
            request.POST,
            instance=stock_doc,
            company=company,
            document_type=doc_type,
        )

        formset = StockDocumentLineFormSet(
            request.POST,
            instance=stock_doc,
            form_kwargs={
                "company": company,
            },
        )

        if form.is_valid() and formset.is_valid():
            valid_lines = 0

            for line_form in formset:
                cleaned = getattr(line_form, "cleaned_data", None)

                if not cleaned:
                    continue

                if cleaned.get("DELETE", False):
                    continue

                item = cleaned.get("item")
                qty = cleaned.get("qty") or Decimal("0.00")
                unit_cost = cleaned.get("unit_cost") or Decimal("0.00")

                if item and qty > 0 and unit_cost >= 0:
                    valid_lines += 1

            if valid_lines < 1:
                messages.error(request, "Please input at least one item line.")
            else:
                with transaction.atomic():
                    doc = form.save(commit=False)
                    doc.company = company
                    doc.document_type = doc_type
                    doc.created_by = request.user
                    doc.save()

                    formset.instance = doc
                    formset.save()

                    create_stock_journal(doc, request.user)

                messages.success(request, "Stock document saved successfully.")
                return redirect("stock_document_list", doc_type=doc_type)

    else:
        form = StockDocumentForm(
            instance=stock_doc,
            company=company,
            document_type=doc_type,
            initial={
                "document_date": timezone.localdate(),
                "status": StockDocument.STATUS_POSTED,
            },
        )

        formset = StockDocumentLineFormSet(
            instance=stock_doc,
            form_kwargs={
                "company": company,
            },
        )

    title_map = {
        StockDocument.TYPE_ISSUE: "Create Stock Issue",
        StockDocument.TYPE_ADJUSTMENT: "Create Stock Adjustment",
        StockDocument.TYPE_ASSEMBLY: "Create Stock Assembly",
        StockDocument.TYPE_TRANSFER: "Create Stock Transfer",
    }

    return render(request, "stock/stock_document_form.html", {
        "company": company,
        "form": form,
        "formset": formset,
        "doc_type": doc_type,
        "page_title": title_map.get(doc_type, "Create Stock Document"),
        "button_text": "Save & Generate Journal",
    })


# =========================================================
# SHORTCUT LIST VIEWS
# =========================================================

@login_required
def stock_issue_list(request):
    return stock_document_list(
        request,
        StockDocument.TYPE_ISSUE,
    )


@login_required
def stock_adjustment_list(request):
    return stock_document_list(
        request,
        StockDocument.TYPE_ADJUSTMENT,
    )


@login_required
def stock_assembly_list(request):
    return stock_document_list(
        request,
        StockDocument.TYPE_ASSEMBLY,
    )


@login_required
def stock_transfer_list(request):
    return stock_document_list(
        request,
        StockDocument.TYPE_TRANSFER,
    )


# =========================================================
# SHORTCUT CREATE VIEWS
# =========================================================

@login_required
def stock_issue_create(request):
    return stock_document_create(
        request,
        StockDocument.TYPE_ISSUE,
    )


@login_required
def stock_adjustment_create(request):
    return stock_document_create(
        request,
        StockDocument.TYPE_ADJUSTMENT,
    )


@login_required
def stock_assembly_create(request):
    return stock_document_create(
        request,
        StockDocument.TYPE_ASSEMBLY,
    )


@login_required
def stock_transfer_create(request):
    return stock_document_create(
        request,
        StockDocument.TYPE_TRANSFER,
    )